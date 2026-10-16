"""Validate generated slide decks have sane dimensions and no off-canvas shapes."""
from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1] / "slides"
decks = [
    ROOT / "lecture1_honest_probabilities.pptx",
    ROOT / "lecture2_deciding_under_uncertainty.pptx",
]
for path in decks:
    pres = Presentation(path)
    ratio = pres.slide_width / pres.slide_height
    assert abs(ratio - 16 / 9) < .01, f"Unexpected aspect ratio: {path}"
    assert len(pres.slides) >= 20, f"Not enough slides: {path}"
    slop = 50000  # 0.055 inch tolerance
    for j, slide in enumerate(pres.slides, 1):
        for shape in slide.shapes:
            if (shape.left < -slop or shape.top < -slop or
                shape.left + shape.width > pres.slide_width + slop or
                shape.top + shape.height > pres.slide_height + slop):
                raise AssertionError(f"{path.name}: slide {j}: shape outside canvas")
    print(f"Slide geometry OK: {path.name}, {len(pres.slides)} slides")
